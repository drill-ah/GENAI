from io import BytesIO

from docx import Document
from pypdf import PdfReader


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text from a supported resume file."""
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if suffix == "txt":
        text = content.decode("utf-8-sig")
    elif suffix == "pdf":
        reader = PdfReader(BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == "docx":
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_text = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        text = "\n".join(paragraphs + table_text)
    else:
        raise ValueError("Upload a PDF, DOCX, or TXT resume.")

    if not text.strip():
        raise ValueError(
            "No readable text was found. If this is a scanned PDF, run OCR "
            "before uploading it."
        )

    return text.strip()
